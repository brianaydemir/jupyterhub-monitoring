"""Tests for report delivery to stdout, files, and email."""

import datetime as dt
import io
import smtplib
from email.header import decode_header, make_header
from email.utils import parseaddr
from pathlib import Path
from typing import Any
from zoneinfo import ZoneInfo

import pytest
from openpyxl import load_workbook

from app.reports import email_utils
from app.reports.builders import build_activity_report, build_new_users_report
from app.reports.delivery import OutputTargets, deliver_report
from app.reports.email_utils import EmailConfig, create_message, send_report_email
from app.reports.model import ReportAttachment
from app.reports.renderers import (
    render_csv_bytes,
    render_html_document,
    render_html_fragment,
    render_text,
    render_xlsx_bytes,
)
from app.usage.activity import ActivitySummary, SessionActivity
from tests.helpers import sample_reports

CHICAGO = ZoneInfo("America/Chicago")
START = dt.datetime(2026, 7, 1, tzinfo=CHICAGO)
END = dt.datetime(2026, 7, 8, tzinfo=CHICAGO)


class _RecordingSmtp:
    """Captures the message a send would have transmitted."""

    sent: list[tuple[str, list[str], Any]] = []

    def __init__(self, host: str, port: int, timeout: float) -> None:
        """Record the connection target."""
        self.host = host
        self.port = port
        self.timeout = timeout

    def __enter__(self) -> "_RecordingSmtp":
        """Support use as a context manager, like :mod:`smtplib` does."""
        return self

    def __exit__(self, *_: object) -> None:
        """Nothing to release."""

    def send_message(self, message: Any, sender: str, recipients: list[str]) -> None:
        """Record the message rather than transmitting it."""
        type(self).sent.append((sender, recipients, message))


def test_deliver_report_prints_to_stdout(capsys: pytest.CaptureFixture[str]) -> None:
    """The report always goes to stdout, whatever else is configured."""
    assert deliver_report(_targets(), sample_reports.report()) == 0
    assert "Total active time" in capsys.readouterr().out


def test_deliver_report_writes_every_requested_file(tmp_path: Path) -> None:
    """Each output flag writes its own format."""
    built = sample_reports.report()
    targets = _targets(
        text_file=tmp_path / "r.txt",
        html_file=tmp_path / "r.html",
        csv_file=tmp_path / "r.csv",
        xlsx_file=tmp_path / "r.xlsx",
    )
    assert deliver_report(targets, built) == 0

    assert (tmp_path / "r.txt").read_text(encoding="utf-8") == render_text(built) + "\n"
    assert (tmp_path / "r.csv").read_bytes() == render_csv_bytes(built)
    assert load_workbook(io.BytesIO((tmp_path / "r.xlsx").read_bytes())).sheetnames


def test_the_html_file_is_a_standalone_document(tmp_path: Path) -> None:
    """A file is opened by a browser, so it gets the whole document."""
    written = tmp_path / "r.html"
    built = sample_reports.report()
    assert deliver_report(_targets(html_file=written), built) == 0
    assert written.read_text(encoding="utf-8") == render_html_document(built)


def test_deliver_report_names_each_file_it_wrote(
    tmp_path: Path,
    capsys: pytest.CaptureFixture[str],
) -> None:
    """Every written path is reported, on stderr so stdout stays the report."""
    targets = _targets(
        text_file=tmp_path / "r.txt",
        html_file=tmp_path / "r.html",
        csv_file=tmp_path / "r.csv",
        xlsx_file=tmp_path / "r.xlsx",
    )
    deliver_report(targets, sample_reports.report())
    err = capsys.readouterr().err
    for suffix in ("txt", "html", "csv", "xlsx"):
        assert str(tmp_path / f"r.{suffix}") in err


def test_deliver_report_sends_email_with_the_table_attached(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """The email carries both bodies and the flagged table as CSV."""
    sent = _capture_smtp(monkeypatch)
    built = sample_reports.report()
    assert deliver_report(_targets(send_email=True), built) == 0

    sender, recipients, message = sent[0]
    assert (sender, recipients) == ("from@example.edu", ["to@example.org"])
    payloads = _payload_names(message)
    assert payloads == ["text/plain", "text/html", "summary.csv"]


def test_the_workbook_is_attached_only_when_it_was_written(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    """``--xlsx-file`` produces the workbook, and the email carries that file."""
    sent = _capture_smtp(monkeypatch)
    built = sample_reports.report()

    deliver_report(_targets(send_email=True), built)
    assert "r.xlsx" not in _payload_names(sent[-1][2])

    written = tmp_path / "r.xlsx"
    deliver_report(_targets(send_email=True, xlsx_file=written), built)
    attached = [
        part.get_payload(decode=True)
        for part in sent[-1][2].walk()
        if part.get_filename() == "r.xlsx"
    ]
    assert attached == [written.read_bytes()]


@pytest.mark.parametrize(
    "failure",
    [smtplib.SMTPConnectError(421, "refused"), TimeoutError("timed out")],
)
def test_deliver_report_reports_a_failed_send(
    failure: Exception,
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
    capsys: pytest.CaptureFixture[str],
) -> None:
    """A refused or timed-out send is an exit code, not a traceback.

    The files were written before the send was attempted, so they are
    not lost to it.
    """

    def _fail(host: str, port: int, timeout: float) -> Any:
        raise failure

    monkeypatch.setattr(email_utils.smtplib, "SMTP_SSL", _fail)
    written = tmp_path / "r.txt"
    targets = _targets(text_file=written, send_email=True)
    assert deliver_report(targets, sample_reports.report()) == 1
    assert "Error sending email" in capsys.readouterr().err
    assert written.exists()


def test_email_uses_ssl_unless_it_is_disabled(monkeypatch: pytest.MonkeyPatch) -> None:
    """SSL is the default; ``--smtp-no-ssl`` selects the plain transport."""
    used: list[str] = []

    def _record(kind: str) -> Any:
        def _factory(host: str, port: int, timeout: float) -> _RecordingSmtp:
            used.append(kind)
            return _RecordingSmtp(host, port, timeout)

        return _factory

    monkeypatch.setattr(email_utils.smtplib, "SMTP_SSL", _record("ssl"))
    monkeypatch.setattr(email_utils.smtplib, "SMTP", _record("plain"))
    deliver_report(_targets(send_email=True), sample_reports.report())
    deliver_report(_targets(send_email=True, smtp_no_ssl=True), sample_reports.report())
    assert used == ["ssl", "plain"]


def test_either_transport_gives_up_on_a_silent_server(monkeypatch: pytest.MonkeyPatch) -> None:
    """Without a timeout, a relay that stops answering hangs a cron job."""
    timeouts: list[float | None] = []

    def _record(host: str, port: int, timeout: float | None = None) -> _RecordingSmtp:
        timeouts.append(timeout)
        return _RecordingSmtp(host, port, timeout or 0)

    monkeypatch.setattr(email_utils.smtplib, "SMTP_SSL", _record)
    monkeypatch.setattr(email_utils.smtplib, "SMTP", _record)
    deliver_report(_targets(send_email=True), sample_reports.report())
    deliver_report(_targets(send_email=True, smtp_no_ssl=True), sample_reports.report())
    assert len(timeouts) == 2
    assert all(timeout is not None and timeout > 0 for timeout in timeouts)


def test_a_message_carries_the_headers_a_relay_expects() -> None:
    """``send_message`` adds neither Date nor Message-ID, so we must."""
    message = create_message(
        sender_name="Reports",
        sender_email="reports@example.edu",
        recipients=["ops@example.org"],
        subject="Subj",
        text_content="t",
        html_content="<p>h</p>",
    )
    assert message["From"] == "Reports <reports@example.edu>"
    assert message["To"] == "ops@example.org"
    assert message["Subject"] == "Subj"
    assert message["Date"]
    assert message["Message-ID"].endswith("@example.edu>")


def test_a_message_omits_a_display_name_when_none_is_given() -> None:
    """A bare address is used as-is rather than wrapped in empty brackets."""
    message = create_message(
        sender_name=None,
        sender_email="reports@example.edu",
        recipients=["ops@example.org"],
        subject="S",
        text_content="t",
        html_content=None,
    )
    assert message["From"] == "reports@example.edu"


def test_a_message_body_survives_non_ascii_content() -> None:
    """Institution domains and identifiers are not all ASCII."""
    message = create_message(
        sender_name=None,
        sender_email="a@example.edu",
        recipients=["b@example.org"],
        subject="S",
        text_content="exämple.edu",
        html_content=None,
    )
    body = [p for p in message.walk() if p.get_content_type() == "text/plain"][0]
    payload = body.get_payload(decode=True)
    assert isinstance(payload, bytes)
    assert payload.decode(str(body.get_content_charset())) == "exämple.edu"


def test_an_attachment_type_is_guessed_from_its_filename() -> None:
    """A hint is optional; the extension is enough for the common cases."""
    message = create_message(
        sender_name=None,
        sender_email="a@example.edu",
        recipients=["b@example.org"],
        subject="S",
        text_content="t",
        html_content=None,
        report_attachments=[ReportAttachment(filename="data.csv", payload=b"a,b\n")],
    )
    part = [p for p in message.walk() if p.get_filename() == "data.csv"][0]
    assert part.get_content_type() == "text/csv"


def _addressed_to(recipients: list[str]) -> Any:
    """Build a message for *recipients*, to assert on its ``To`` header."""
    return create_message(
        sender_name=None,
        sender_email="reports@example.edu",
        recipients=recipients,
        subject="S",
        text_content="t",
        html_content=None,
    )


def test_every_recipient_reaches_the_envelope_and_the_header(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """A list of addresses is a list to smtplib, not one long address."""
    sent = _capture_smtp(monkeypatch)
    config = _email(recipients=("a@example.org", "b@example.org"))
    assert send_report_email(config, text_content="t", html_content="h", attachments=[]) == 0

    _, recipients, message = sent[-1]
    assert recipients == ["a@example.org", "b@example.org"]
    assert message["To"] == "a@example.org, b@example.org"


def test_a_display_name_reaches_the_header_but_not_the_envelope(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """The relay is given addresses; only the header carries the names."""
    sent = _capture_smtp(monkeypatch)
    config = _email(recipients=("Ops <a@example.org>", "b@example.org"))
    assert send_report_email(config, text_content="t", html_content="h", attachments=[]) == 0

    _, recipients, message = sent[-1]
    assert recipients == ["a@example.org", "b@example.org"]
    assert message["To"] == "Ops <a@example.org>, b@example.org"


def test_one_entry_may_carry_several_mailboxes(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """Entries are parsed, not counted, so a list in one reaches everyone.

    ``--recipient-email`` accepts a list in a single occurrence, which
    arrives here as one entry holding two mailboxes.
    """
    sent = _capture_smtp(monkeypatch)
    config = _email(recipients=("Ops <a@example.org>, b@example.org",))
    assert send_report_email(config, text_content="t", html_content="h", attachments=[]) == 0

    _, recipients, message = sent[-1]
    assert recipients == ["a@example.org", "b@example.org"]
    assert message["To"] == "Ops <a@example.org>, b@example.org"


def test_a_quoted_comma_in_a_list_does_not_separate_mailboxes() -> None:
    """The entry is parsed as a header field, so the quotes are honored."""
    message = _addressed_to(['"Doe, Jane" <j@example.org>, dean@example.org'])
    assert message["To"] == '"Doe, Jane" <j@example.org>, dean@example.org'


def test_a_quoted_display_name_keeps_its_quotes_and_its_address(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """The comma inside the name must not read as the end of the mailbox."""
    sent = _capture_smtp(monkeypatch)
    config = _email(recipients=('"Doe, Jane" <jane@example.org>',))
    assert send_report_email(config, text_content="t", html_content="h", attachments=[]) == 0

    _, recipients, message = sent[-1]
    assert recipients == ["jane@example.org"]
    assert message["To"] == '"Doe, Jane" <jane@example.org>'


def test_a_non_ascii_display_name_is_encoded_but_its_address_is_not() -> None:
    """Encoding the mailbox whole would bury the address in the word.

    The choice between base64 and quoted-printable is the shortest of
    the two for the name at hand, so the test asserts what has to hold
    of either: the header is ASCII, the address is still readable in
    it, and the name decodes back to what was given.
    """
    message = _addressed_to(["exämple Ops <ops@example.org>"])
    header = message["To"]
    assert header.isascii()
    assert header.endswith(" <ops@example.org>")
    assert b" <ops@example.org>" in message.as_bytes()
    name, address = parseaddr(header)
    assert address == "ops@example.org"
    assert str(make_header(decode_header(name))) == "exämple Ops"


def test_one_unreadable_recipient_does_not_cost_the_others_their_copy(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """Parsing the entries as one field would drop every one of them.

    An unbalanced quote passes the CLI's per-entry check, so it can
    reach here; parsing each entry on its own keeps the damage to the
    entry that caused it.
    """
    sent = _capture_smtp(monkeypatch)
    config = _email(recipients=('"Doe <a@example.org>', "b@example.org"))
    assert send_report_email(config, text_content="t", html_content="h", attachments=[]) == 0

    _, recipients, message = sent[-1]
    assert "b@example.org" in recipients
    assert "b@example.org" in message["To"]


def test_an_anonymized_report_leaks_no_identifier_in_any_format() -> None:
    """The pseudonyms exist to keep raw identifiers out of every output."""
    names = ["alice@x.edu", "eppn:alice@x.edu", "orcid:0009-0008-3064-0494", "svc-account"]
    activity = build_activity_report(
        sessions=[
            SessionActivity(
                user=name,
                server_name="",
                started="2026-07-02T10:00:00Z",
                seconds=3600,
                gpu=False,
            )
            for name in names
        ],
        summary=ActivitySummary(total_seconds=14400, per_day_seconds={"2026-07-02": 14400}),
        start_time=START,
        end_time=END,
        strftime_fmt="%Y-%m-%d %H:%M",
        anonymize=True,
    )
    users = build_new_users_report(
        users=[{"name": name, "created": "2026-07-02T10:00:00Z"} for name in names],
        start_time=START,
        end_time=END,
        strftime_fmt="%Y-%m-%d",
        anonymize=True,
        first_server={name: 1782000000 for name in names},
    )
    secrets = ["alice", "0009-0008-3064-0494", "svc-account"]
    for built in (activity, users):
        for rendered in (
            render_text(built),
            render_html_fragment(built),
            render_csv_bytes(built).decode("utf-8"),
        ):
            assert not [s for s in secrets if s in rendered]
        # The XLSX bytes are compressed, so the cells are read back instead.
        workbook = load_workbook(io.BytesIO(render_xlsx_bytes(built)))
        cells = [
            str(value)
            for sheet in workbook.worksheets
            for row in sheet.iter_rows(values_only=True)
            for value in row
        ]
        assert not [s for s in secrets for cell in cells if s in cell]
        # The institution is deliberately preserved.
        assert "x.edu" in render_text(built)


def _targets(
    *,
    text_file: Path | None = None,
    html_file: Path | None = None,
    csv_file: Path | None = None,
    xlsx_file: Path | None = None,
    send_email: bool = False,
    smtp_no_ssl: bool = False,
) -> OutputTargets:
    """Build the delivery targets a report command would hand to delivery."""
    return OutputTargets(
        text_file=text_file,
        html_file=html_file,
        csv_file=csv_file,
        xlsx_file=xlsx_file,
        email=_email(smtp_no_ssl=smtp_no_ssl) if send_email else None,
    )


def _email(
    *,
    recipients: tuple[str, ...] = ("to@example.org",),
    smtp_no_ssl: bool = False,
) -> EmailConfig:
    """Build the email settings the delivery tests send with."""
    return EmailConfig(
        sender_name=None,
        sender_email="from@example.edu",
        recipients=recipients,
        subject="Test Report",
        smtp_host="smtp.example.edu",
        smtp_port=465,
        smtp_no_ssl=smtp_no_ssl,
    )


def _capture_smtp(monkeypatch: pytest.MonkeyPatch) -> list[tuple[str, list[str], Any]]:
    """Replace both SMTP transports with a recorder and return its log."""
    _RecordingSmtp.sent = []
    monkeypatch.setattr(email_utils.smtplib, "SMTP_SSL", _RecordingSmtp)
    monkeypatch.setattr(email_utils.smtplib, "SMTP", _RecordingSmtp)
    return _RecordingSmtp.sent


def _payload_names(message: Any) -> list[str]:
    """Return each leaf part of *message* by filename, or by content type."""
    return [
        part.get_filename() or part.get_content_type()
        for part in message.walk()
        if not part.is_multipart()
    ]
