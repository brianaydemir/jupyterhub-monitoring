"""Render report-domain models into plain text, HTML, CSV, and XLSX.

Plain text and HTML return ``str``,
because a caller goes on to print or embed them;
CSV and XLSX return ``bytes``,
because they only ever reach a file or an email attachment,
and their names say so.

Text, CSV, and HTML show a cell as :func:`~app.reports.model.cell_text`
renders it, though CSV escapes text a spreadsheet would otherwise open
as a formula; see :func:`_csv_cell`.
XLSX alone writes the underlying value,
so that a workbook holds numbers to sort, sum, and chart
rather than the text of them.
"""

import csv
import io
import re
from collections.abc import Sequence
from dataclasses import dataclass
from typing import ClassVar, Literal

from jinja2 import BaseLoader, Environment
from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet
from tabulate import tabulate

from app.reports.model import (
    Cell,
    Duration,
    Fixed,
    Percent,
    Report,
    ReportAttachment,
    ReportSection,
    SectionBlock,
    TableBlock,
    TextBlock,
    cell_text,
)

_ENV = Environment(loader=BaseLoader(), autoescape=True, trim_blocks=True, lstrip_blocks=True)

# Footnote headings, shared by every renderer
# so the four output formats label the same content identically.
_SECTION_NOTES_LABEL = "Notes"
_REPORT_NOTES_LABEL = "Report notes"

# Name of the leading XLSX sheet holding the report's narrative.
_INFO_SHEET_NAME = "Report"

# Excel number formats for the typed cells,
# and the unit a duration is written in:
# Excel measures elapsed time in days, and "[h]:mm" is the format
# that shows such a value as hours past 24 rather than wrapping.
_SECONDS_PER_DAY = 24 * 60 * 60
_DURATION_FORMAT = "[h]:mm"
_PERCENT_FORMAT = "0%"
_COUNT_FORMAT = "General"

# The classes a rendered table's cells carry.
# The template selects the dark-mode rules on them,
# and :func:`_cell_style` names the one an alternating row gets,
# so both sides have to agree on the spelling.
_HEADER_CLASS = "th"
_ROW_CLASS = "td"
_ROW_ALT_CLASS = "td-alt"

# Padding a table cell carries in either color scheme.
_CELL_PADDING = "2px 8px"


@dataclass(frozen=True)
class _TablePalette:
    """The colors one color scheme gives a table."""

    header_background: str
    header_text: str
    row_background: str
    row_alt_background: str
    row_text: str
    border: str


#: Written inline on every cell, because an email client may strip a stylesheet
#: but keeps a ``style`` attribute.
_LIGHT_PALETTE = _TablePalette(
    header_background="#a6c9e8",
    header_text="#000000",
    row_background="#ffffff",
    row_alt_background="#e6eff4",
    row_text="#000000",
    border="#9ab3c8",
)

#: Written to the template's ``<style>`` block,
#: since only a media query can ask which color scheme the reader is in.
_DARK_PALETTE = _TablePalette(
    header_background="#1d4060",
    header_text="#d0e8f8",
    row_background="#18232d",
    row_alt_background="#1e2f3d",
    row_text="#d8e8f0",
    border="#3a607a",
)

_HTML_TEMPLATE = _ENV.from_string("""
<style>
  @media (prefers-color-scheme: dark) {
    {# Every declaration below needs !important: the light-mode colors
       are inline on each cell, and an inline style outranks a
       stylesheet rule. Drop it and dark mode silently stops applying. #}
    .{{ header_class }} {
      background: {{ dark_palette.header_background }} !important;
      color: {{ dark_palette.header_text }} !important;
      border-color: {{ dark_palette.border }} !important;
    }
    .{{ row_class }} {
      background: {{ dark_palette.row_background }} !important;
      color: {{ dark_palette.row_text }} !important;
      border-color: {{ dark_palette.border }} !important;
    }
    .{{ row_alt_class }} {
      background: {{ dark_palette.row_alt_background }} !important;
      color: {{ dark_palette.row_text }} !important;
      border-color: {{ dark_palette.border }} !important;
    }
  }
</style>
<h2>{{ report.title }}</h2>
{% for section in report.sections %}
  <h3>{{ section.heading }}</h3>
  <p>{{ section.description }}</p>
  {% for block in section.blocks %}
    {% if block.kind == "table" %}
      <table style="border-collapse:collapse">
        <thead>
          <tr>
          {% for header in block.headers %}
            <th class="{{ header_class }}" style="{{ header.style }}">{{ header.text }}</th>
          {% endfor %}
          </tr>
        </thead>
        <tbody>
        {% for row in block.rows %}
          <tr>
          {% for cell in row %}
            <td class="{{ cell.cls }}" style="{{ cell.style }}">{{ cell.text }}</td>
          {% endfor %}
          </tr>
        {% endfor %}
        </tbody>
      </table>
    {% else %}
      <p>{{ block.text }}</p>
    {% endif %}
  {% endfor %}
  {% if section.footnotes %}
    <h4>{{ section_notes_label }}</h4>
    <ul>
    {% for footnote in section.footnotes %}
      <li><em>{{ footnote }}</em></li>
    {% endfor %}
    </ul>
  {% endif %}
{% endfor %}
{% if report.footnotes %}
  <h4>{{ report_notes_label }}</h4>
  <ul>
  {% for footnote in report.footnotes %}
    <li><em>{{ footnote }}</em></li>
  {% endfor %}
  </ul>
{% endif %}
""")


_HTML_DOCUMENT_TEMPLATE = _ENV.from_string("""<!doctype html>
<html lang="en">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>{{ title }}</title>
</head>
<body>
{{ body | safe }}
</body>
</html>
""")


def _display_rows(block: TableBlock) -> list[list[str]]:
    """Render a table block's cells to the strings the plain-text table shows."""
    return [[cell_text(value) for value in row] for row in block.rows]


def _render_text_table(block: TableBlock) -> list[str]:
    """Render a table block as aligned plain-text lines.

    ``colalign`` is omitted for a table with no rows,
    because tabulate indexes it against the data rows
    and fails when there are none.
    """
    colalign = (
        tuple(
            "right" if idx in block.right_align_columns else "left"
            for idx, _ in enumerate(block.headers)
        )
        if block.rows
        else None
    )
    table = tabulate(
        _display_rows(block),
        headers=block.headers,
        tablefmt="simple",
        colalign=colalign,
        disable_numparse=True,
    )
    return table.splitlines()


def render_text(report: Report) -> str:
    """Render a report as plain text with section headings and tables."""
    lines = [report.title, ""]
    for section in report.sections:
        lines.append(section.heading)
        lines.append("=" * len(section.heading))
        lines.append(section.description)
        lines.append("")
        for block in section.blocks:
            if isinstance(block, TextBlock):
                lines.append(block.text)
                lines.append("")
            else:
                lines.extend(_render_text_table(block))
                lines.append("")
        if section.footnotes:
            lines.append(f"{_SECTION_NOTES_LABEL}:")
            lines.extend(f"- {note}" for note in section.footnotes)
            lines.append("")
    if report.footnotes:
        lines.append(f"{_REPORT_NOTES_LABEL}:")
        lines.extend(f"- {note}" for note in report.footnotes)
    return "\n".join(lines).rstrip()


def _align_value(right_align: bool) -> str:
    """Return the CSS ``text-align`` keyword for a column."""
    return "right" if right_align else "left"


def _inline_style(right_align: bool, background: str, color: str) -> str:
    """Return the inline CSS every table cell carries, header or body."""
    return (
        f"text-align:{_align_value(right_align)}; "
        f"border:1px solid {_LIGHT_PALETTE.border}; padding:{_CELL_PADDING}; "
        f"background:{background}; color:{color}"
    )


def _header_style(right_align: bool) -> str:
    """Return the inline CSS for a table header cell."""
    return _inline_style(
        right_align,
        background=_LIGHT_PALETTE.header_background,
        color=_LIGHT_PALETTE.header_text,
    )


def _cell_style(row_idx: int, right_align: bool) -> tuple[str, str]:
    """Return the CSS class and inline style for a body cell."""
    alternate = bool(row_idx % 2)
    cls = _ROW_ALT_CLASS if alternate else _ROW_CLASS
    background = (
        _LIGHT_PALETTE.row_alt_background if alternate else _LIGHT_PALETTE.row_background
    )
    style = _inline_style(right_align, background=background, color=_LIGHT_PALETTE.row_text)
    return cls, style


@dataclass(frozen=True)
class _HeaderCell:
    """A header cell with its inline CSS resolved."""

    text: str
    style: str


@dataclass(frozen=True)
class _BodyCell:
    """A body cell with its alternating-row class and inline CSS resolved."""

    text: str
    cls: str
    style: str


@dataclass(frozen=True)
class _ShapedText:
    """A text block as the template sees it."""

    kind: ClassVar[Literal["text"]] = "text"
    text: str


@dataclass(frozen=True)
class _ShapedTable:
    """A table block as the template sees it."""

    kind: ClassVar[Literal["table"]] = "table"
    headers: list[_HeaderCell]
    rows: list[list[_BodyCell]]


#: The template's counterpart to :data:`~app.reports.model.SectionBlock`,
#: tagged by ``kind`` because Jinja cannot narrow a union itself.
_ShapedBlock = _ShapedText | _ShapedTable


@dataclass(frozen=True)
class _ShapedSection:
    """A section whose blocks have been shaped for the template."""

    heading: str
    description: str
    blocks: list[_ShapedBlock]
    footnotes: list[str]


@dataclass(frozen=True)
class _ShapedReport:
    """The whole context the HTML template renders from."""

    title: str
    sections: list[_ShapedSection]
    footnotes: list[str]


def _shape_table_block(block: TableBlock) -> _ShapedTable:
    """Resolve a table block's classes and inline styles."""
    headers = [
        _HeaderCell(text=header, style=_header_style(idx in block.right_align_columns))
        for idx, header in enumerate(block.headers)
    ]
    rows: list[list[_BodyCell]] = []
    for row_idx, row in enumerate(block.rows):
        cells: list[_BodyCell] = []
        for col_idx, value in enumerate(row):
            cls, style = _cell_style(row_idx, col_idx in block.right_align_columns)
            cells.append(_BodyCell(text=cell_text(value), cls=cls, style=style))
        rows.append(cells)
    return _ShapedTable(headers=headers, rows=rows)


def _shape_block(block: SectionBlock) -> _ShapedBlock:
    """Shape one block, preserving the model's text/table distinction."""
    if isinstance(block, TextBlock):
        return _ShapedText(text=block.text)
    return _shape_table_block(block)


def _shape_report(report: Report) -> _ShapedReport:
    """Build the template context, so the template does no styling itself."""
    return _ShapedReport(
        title=report.title,
        sections=[
            _ShapedSection(
                heading=section.heading,
                description=section.description,
                blocks=[_shape_block(block) for block in section.blocks],
                footnotes=section.footnotes,
            )
            for section in report.sections
        ],
        footnotes=report.footnotes,
    )


def render_html_fragment(report: Report) -> str:
    """Render a report as an HTML fragment, suitable for an email body."""
    return _HTML_TEMPLATE.render(
        report=_shape_report(report),
        section_notes_label=_SECTION_NOTES_LABEL,
        report_notes_label=_REPORT_NOTES_LABEL,
        header_class=_HEADER_CLASS,
        row_class=_ROW_CLASS,
        row_alt_class=_ROW_ALT_CLASS,
        dark_palette=_DARK_PALETTE,
    )


def render_html_document(report: Report) -> str:
    """Render a report as a standalone HTML document.

    :func:`render_html_fragment` produces markup for an email body;
    a file a browser opens directly
    also needs a doctype and a charset declaration.
    """
    return _HTML_DOCUMENT_TEMPLATE.render(
        title=report.title,
        body=render_html_fragment(report),
    )


def report_attachments(report: Report) -> list[ReportAttachment]:
    """Collect the report's attachments, de-duplicated by filename.

    Each table block carrying an ``attachment_filename`` is rendered to CSV.
    Pre-built ``ReportSection.attachments`` are collected too,
    though no builder populates them yet.
    On a filename collision the last attachment encountered wins.
    """
    attachments: list[ReportAttachment] = []
    for section in report.sections:
        attachments.extend(section.attachments)
        for block in section.blocks:
            if isinstance(block, TableBlock) and block.attachment_filename:
                attachments.append(
                    ReportAttachment(
                        filename=block.attachment_filename,
                        payload=_table_to_csv_bytes(block),
                        mime_hint="text/csv",
                    )
                )

    deduped: dict[str, ReportAttachment] = {}
    for attachment in attachments:
        deduped[attachment.filename] = attachment
    return list(deduped.values())


# The characters Excel and LibreOffice read as the start of a formula
# when they open a CSV.
# A leading apostrophe is the conventional neutralizer:
# the spreadsheet drops it and shows the rest as text.
_CSV_FORMULA_LEAD = ("=", "+", "-", "@", "\t", "\r")


def _csv_cell(value: Cell) -> str:
    """Return a cell's CSV text, neutralizing a formula-like value.

    Only text is escaped, for the reason
    :func:`_append_row` forces a text cell to the text type:
    report cells carry text
    derived from JupyterHub usernames and named-server names.
    A number is written as the other text formats show it,
    so that a negative duration keeps reading as ``-0:30``
    rather than as an escaped string.
    """
    if isinstance(value, str):
        return f"'{value}" if value.startswith(_CSV_FORMULA_LEAD) else value
    return cell_text(value)


def _csv_row(values: Sequence[Cell]) -> list[str]:
    """Render one CSV row, so that the two CSV entry points cannot drift."""
    return [_csv_cell(value) for value in values]


def _table_to_csv_bytes(block: TableBlock) -> bytes:
    """Render a single table block as UTF-8 CSV bytes."""
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(_csv_row(block.headers))
    writer.writerows(_csv_row(row) for row in block.rows)
    return buf.getvalue().encode("utf-8")


def render_csv_bytes(report: Report) -> bytes:
    """Render a report to a single UTF-8 CSV document.

    Every section is written to one document, separated by blank rows.
    """
    buf = io.StringIO()
    writer = csv.writer(buf)

    writer.writerow(_csv_row([report.title]))
    writer.writerow([])

    for section in report.sections:
        writer.writerow(_csv_row([section.heading]))
        writer.writerow(_csv_row([section.description]))
        writer.writerow([])

        for block in section.blocks:
            if isinstance(block, TextBlock):
                writer.writerow(_csv_row([block.text]))
                writer.writerow([])
                continue

            writer.writerow(_csv_row(block.headers))
            writer.writerows(_csv_row(row) for row in block.rows)
            writer.writerow([])

        if section.footnotes:
            writer.writerow(_csv_row([_SECTION_NOTES_LABEL]))
            for note in section.footnotes:
                writer.writerow(_csv_row([note]))
            writer.writerow([])

    if report.footnotes:
        writer.writerow(_csv_row([_REPORT_NOTES_LABEL]))
        for note in report.footnotes:
            writer.writerow(_csv_row([note]))

    return buf.getvalue().encode("utf-8")


def render_xlsx_bytes(report: Report) -> bytes:
    """Render a report to an XLSX workbook, one sheet per table block.

    A leading "Report" sheet carries the title and each section's heading,
    description, text blocks, and footnotes.
    Without it the workbook would be the only output format recording
    neither the reporting window nor the timezone the timestamps are in.

    Each table sheet then repeats its own section's footnotes below the table.
    The duplication is intended: a reader who clicks straight to a data sheet
    never sees the "Report" sheet, which stays the narrative overview
    and remains the only place a section holding no table,
    or holding a text block, has to state its notes.

    Raises:
        ValueError: If openpyxl does not provide an initial worksheet.
    """
    workbook = Workbook()
    info_sheet = workbook.active
    if info_sheet is None:
        raise ValueError("Workbook did not create an initial worksheet")
    info_sheet.title = _INFO_SHEET_NAME

    # Seeded with the info sheet,
    # so a section of the same name is given a distinct sheet here
    # rather than being renamed by openpyxl alone.
    used_sheet_names: set[str] = {_INFO_SHEET_NAME}
    table_count = 0

    _append_row(info_sheet, [report.title])

    for section in report.sections:
        table_blocks = _append_section_summary(info_sheet, section)
        table_count += len(table_blocks)
        for idx, block in enumerate(table_blocks, 1):
            title = _table_sheet_title(section.heading, idx, len(table_blocks))
            worksheet = workbook.create_sheet(title=_unique_sheet_name(title, used_sheet_names))
            _append_row(worksheet, block.headers)
            for row in block.rows:
                _append_row(worksheet, row)
            _append_section_notes(worksheet, section.footnotes)

    if report.footnotes:
        info_sheet.append([])
        _append_row(info_sheet, [_REPORT_NOTES_LABEL])
        for footnote in report.footnotes:
            _append_row(info_sheet, [footnote])

    if table_count == 0:
        info_sheet.append([])
        _append_row(info_sheet, ["No table data available."])

    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def _append_section_summary(worksheet: Worksheet, section: ReportSection) -> list[TableBlock]:
    """Write a section's narrative to the info sheet.

    Returns:
        The section's table blocks, which get sheets of their own.
    """
    worksheet.append([])
    _append_row(worksheet, [section.heading])
    _append_row(worksheet, [section.description])

    table_blocks: list[TableBlock] = []
    for block in section.blocks:
        if isinstance(block, TextBlock):
            _append_row(worksheet, [block.text])
        else:
            table_blocks.append(block)

    if section.footnotes:
        _append_row(worksheet, [_SECTION_NOTES_LABEL])
        for footnote in section.footnotes:
            _append_row(worksheet, [footnote])

    return table_blocks


def _append_section_notes(worksheet: Worksheet, footnotes: Sequence[str]) -> None:
    """Repeat a section's footnotes beneath the table on its own sheet.

    Every sheet of a multi-table section gets the same notes,
    since the notes qualify the section rather than any one of its tables.
    The blank row separates them from the data,
    so that a note neither reads as a row of the table
    nor comes along with the table's range.
    """
    if not footnotes:
        return

    worksheet.append([])
    _append_row(worksheet, [_SECTION_NOTES_LABEL])
    for footnote in footnotes:
        _append_row(worksheet, [footnote])


def _table_sheet_title(heading: str, index: int, table_count: int) -> str:
    """Return the sheet title for a section's *index*-th table block.

    The index is appended only when the section has more than one table,
    so the common single-table case keeps the bare section heading.
    """
    if table_count > 1:
        return f"{heading} {index}"
    return heading


def _fixed_format(digits: int) -> str:
    """Return the Excel format showing a number to *digits* decimal places."""
    return "0." + "0" * digits


def _xlsx_cell(value: Cell) -> tuple[str | float, str | None]:
    """Return a cell's workbook value and the number format it reads in.

    A duration becomes the day fraction ``[h]:mm`` displays,
    and a percentage the 0..1 fraction ``0%`` displays,
    so that a column of them sorts and sums as the quantity it is
    rather than as text, in which ``"10:00"`` precedes ``"2:00"``.

    Returns:
        The value, and the number format for a number or ``None`` for text,
        which :func:`_append_row` writes as text instead.
    """
    if isinstance(value, str):
        return value, None
    if isinstance(value, Duration):
        return value.seconds / _SECONDS_PER_DAY, _DURATION_FORMAT
    if isinstance(value, Percent):
        return value.fraction, _PERCENT_FORMAT
    if isinstance(value, Fixed):
        return value.value, _fixed_format(value.digits)
    return value, _COUNT_FORMAT


def _append_row(worksheet: Worksheet, values: Sequence[Cell]) -> None:
    """Append a row, writing each cell as the kind of value it holds.

    A text cell is forced to the text type,
    because openpyxl infers a formula from a leading ``=``
    and report cells carry text derived from JupyterHub usernames.
    A number keeps its own type, which is what the workbook is for,
    and no number can be read as a formula.
    """
    resolved = [_xlsx_cell(value) for value in values]
    worksheet.append([value for value, _ in resolved])
    row = worksheet.max_row
    for column, (_, number_format) in enumerate(resolved, 1):
        cell = worksheet.cell(row=row, column=column)
        if number_format is None:
            cell.data_type = "s"
        else:
            cell.number_format = number_format


def _sanitize_sheet_name(name: str) -> str:
    """Strip the characters Excel rejects in a worksheet name."""
    cleaned = re.sub(r"[\[\]:*?/\\]", "_", name).strip().strip("'")
    return cleaned or "Sheet"


def _unique_sheet_name(base: str, used_names: set[str]) -> str:
    """Return a sheet name unique within *used_names*, within Excel's limit."""
    base_clean = _sanitize_sheet_name(base)
    candidate = base_clean[:31]

    if candidate not in used_names:
        used_names.add(candidate)
        return candidate

    suffix = 2
    while True:
        suffix_text = f" ({suffix})"
        trimmed = base_clean[: 31 - len(suffix_text)]
        candidate = f"{trimmed}{suffix_text}"
        if candidate not in used_names:
            used_names.add(candidate)
            return candidate
        suffix += 1
